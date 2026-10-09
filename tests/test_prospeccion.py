"""Hoja de comprobación de la skill de prospección."""
import importlib.util
import os

from openpyxl import load_workbook

RUTA = os.path.join(os.path.dirname(__file__), "..", ".claude", "skills", "prospeccion-clinicas",
                    "scripts", "generar_hoja.py")
spec = importlib.util.spec_from_file_location("generar_hoja", RUTA)
gh = importlib.util.module_from_spec(spec)
spec.loader.exec_module(gh)

CSV = """nombre,provincia,especialidad,telefono,horario_huecos
Fisio Uno,Murcia,Fisioterapia,968 111 111,?
Dental Dos,Alicante,Clínica dental,612 345 678,?
Dental Tres,Almería,Dental,950 000 000,Sí
Dental Cuatro,Almería,Dental,,?
"""


def test_valor_y_movil():
    assert gh.valor("Clínica dental") == 2 and gh.valor("Medicina estética") == 2
    assert gh.valor("Fisioterapia") == 1
    assert gh.es_movil("612 34 56 78") == 1 and gh.es_movil("+34 678 499 634") == 1
    assert gh.es_movil("950 681 301") == 0 and gh.es_movil("") == 0


def test_cadena_va_al_final_y_no_es_primera(tmp_path):
    entrada = tmp_path / "c.csv"
    entrada.write_text("nombre,provincia,especialidad,telefono,horario_huecos,web,cadena,notas\n"
                       "Cadena,Almería,Dental,950 1,Sí,?,Sí,grupo\nIndependiente,Almería,Dental,,?,No encontrada,,\n",
                       encoding="utf-8")
    clinicas = gh.leer(str(entrada))
    assert [c["nombre"] for c in clinicas] == ["Independiente", "Cadena"]
    salida = tmp_path / "h.xlsx"
    gh.construir(clinicas, str(salida))
    ws = load_workbook(salida)["Comprobación"]
    assert ws["B2"].value.startswith("★") and not ws["B3"].value
    assert ws["K2"].value == "No encontrada" and ws["P3"].value == "Sí" and ws["Q3"].value == "grupo"
    assert 'IF(K2="Sin web",1,0)' in ws["S2"].value


def test_orden_y_hoja(tmp_path):
    entrada = tmp_path / "c.csv"
    entrada.write_text(CSV, encoding="utf-8")
    clinicas = gh.leer(str(entrada))
    # Dental Tres (3+2) > Dental Dos (1+2+1) > Dental Cuatro y Fisio Uno
    assert [c["nombre"] for c in clinicas][:2] == ["Dental Tres", "Dental Dos"]
    salida = tmp_path / "hoja.xlsx"
    assert gh.construir(clinicas, str(salida)) == 4
    wb = load_workbook(salida)
    ws = wb["Comprobación"]
    assert ws["C2"].value == "Dental Tres" and ws["B2"].value.startswith("★")
    assert "google.com/maps/search" in ws["G2"].hyperlink.target
    assert ws["R2"].value.startswith("=IF(") and ws["T2"].value.startswith("=IF(")
    assert wb["Cómo usarla"]["A1"].value == "Qué es"
