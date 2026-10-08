#!/usr/bin/env python3
"""Genera la hoja de comprobación (xlsx) a partir de una lista de clínicas.

Entrada: CSV con columnas  nombre,provincia,especialidad,telefono,horario_huecos[,web,cadena,notas]
  horario_huecos = Sí | No | ?   (déjalo en ? si no lo has visto)
  web = dirección, «No encontrada» o ?   ·   cadena = Sí si es cadena o franquicia

Uso:
  python generar_hoja.py clinicas.csv prospectos/hoja_comprobacion.xlsx

La hoja lleva fórmulas (prioridad, puntos y ranking) que se recalculan al anotar
lo que se ve en Google Maps. Solo usa openpyxl.
"""
import csv
import re
import sys
from urllib.parse import quote

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

VALOR2 = ("dental", "dentista", "odont", "ortodon", "implant", "estétic", "estetic", "dermat")
PRIMERAS = 10

COLS = ["Orden", "Comprobar primero", "Clínica", "Provincia", "Especialidad", "Teléfono", "Google Maps",
        "Valor especialidad (2 dental/estética, 1 resto)", "Móvil WhatsApp (1 sí, 0 no)",
        "Horario con huecos (mediodía, cierra pronto)",
        "Web propia (de la búsqueda; pon «Sin web» solo si en Maps no hay botón de web)",
        "Nº reseñas en Google", "Queja de teléfono en reseñas recientes", "Cita de la reseña (sin nombre del autor)",
        "Llamada de prueba: no cogieron", "¿Es cadena o franquicia?", "Notas",
        "Prioridad", "Puntos (0-10)", "Ranking"]
WIDTHS = [7, 11, 38, 11, 20, 14, 15, 14, 12, 16, 26, 11, 16, 44, 13, 13, 50, 14, 9, 9]
(ORDEN, PRIM, NOMBRE, PROV, ESP, TEL, MAPS, VALOR, MOVIL, HORARIO, WEB, RESENAS, QUEJA, CITA,
 LLAMADA, CADENA, NOTAS, PRIOR, PUNTOS, RANK) = [chr(64 + i) for i in range(1, 21)]
EDITABLES = (WEB, RESENAS, QUEJA, CITA, LLAMADA, CADENA, NOTAS)
CENTRADAS = (ORDEN, VALOR, MOVIL, HORARIO, RESENAS, QUEJA, LLAMADA, CADENA, PRIOR, PUNTOS, RANK)
F = "Arial"


def font(**k):
    return Font(name=F, size=10, **k)


def valor(especialidad):
    t = especialidad.lower()
    return 2 if any(k in t for k in VALOR2) else 1


def es_movil(tel):
    d = re.sub(r"\D", "", tel or "")
    d = d[2:] if d.startswith("34") and len(d) > 9 else d
    return 1 if len(d) == 9 and d[0] in "67" else 0


def puntos_iniciales(c):
    if c.get("cadena") == "Sí":
        return -1
    base = 3 if c["horario_huecos"] == "Sí" else 1
    return base + valor(c["especialidad"]) + es_movil(c["telefono"])


def leer(ruta):
    with open(ruta, newline="", encoding="utf-8") as f:
        filas = [{k.strip(): (v or "").strip() for k, v in r.items()} for r in csv.DictReader(f)]
    for c in filas:
        c["horario_huecos"] = c.get("horario_huecos") or "?"
        c["web"] = c.get("web") or "?"
    # Más puntos primero; a igualdad, las que tienen teléfono publicado.
    filas.sort(key=lambda c: (-puntos_iniciales(c), not c.get("telefono")))
    return filas


def construir(clinicas, salida):
    wb = Workbook()
    ws = wb.active
    ws.title = "Comprobación"
    head = PatternFill("solid", fgColor="1F4E3D")
    amarillo = PatternFill("solid", fgColor="FFF2B3")
    azul = PatternFill("solid", fgColor="E8F0FE")
    verde = PatternFill("solid", fgColor="DDEFE3")
    lado = Side(style="thin", color="D0D0D0")
    borde = Border(left=lado, right=lado, top=lado, bottom=lado)

    ws.append(COLS)
    for c in range(1, len(COLS) + 1):
        x = ws.cell(1, c)
        x.font, x.fill, x.border = font(bold=True, color="FFFFFF"), head, borde
        x.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[1].height = 78

    n = len(clinicas)
    primero, ultimo = 2, 1 + n
    col = lambda letra: ord(letra) - 64
    for i, c in enumerate(clinicas):
        r = 2 + i
        ws[f"{ORDEN}{r}"] = i + 1
        ws[f"{PRIM}{r}"] = "★ Primero" if i < PRIMERAS and c.get("cadena") != "Sí" else ""
        ws[f"{NOMBRE}{r}"] = c["nombre"]
        ws[f"{PROV}{r}"] = c["provincia"]
        ws[f"{ESP}{r}"] = c["especialidad"]
        ws[f"{TEL}{r}"] = c.get("telefono") or "—"
        ws[f"{MAPS}{r}"] = "Abrir en Maps"
        ws[f"{MAPS}{r}"].hyperlink = (
            "https://www.google.com/maps/search/?api=1&query=" + quote(f"{c['nombre']} {c['provincia']}"))
        ws[f"{VALOR}{r}"] = valor(c["especialidad"])
        ws[f"{MOVIL}{r}"] = es_movil(c.get("telefono"))
        ws[f"{HORARIO}{r}"] = c["horario_huecos"]
        ws[f"{WEB}{r}"] = c["web"]
        ws[f"{CADENA}{r}"] = c.get("cadena") or None
        ws[f"{NOTAS}{r}"] = c.get("notas") or None
        ws[f"{PRIOR}{r}"] = (f'=IF({CADENA}{r}="Sí","Descartar",IF(OR({QUEJA}{r}="Sí",{LLAMADA}{r}="Sí"),"Alta",'
                             f'IF({HORARIO}{r}="Sí","Media","Por verificar")))')
        ws[f"{PUNTOS}{r}"] = (f'=IF({CADENA}{r}="Sí",0,MIN(10,IF({PRIOR}{r}="Alta",5,IF({PRIOR}{r}="Media",3,1))'
                              f'+{VALOR}{r}+{MOVIL}{r}+IF({RESENAS}{r}>100,1,0)+IF({WEB}{r}="Sin web",1,0)))')
        ws[f"{RANK}{r}"] = f'=IF({CADENA}{r}="Sí","",RANK({PUNTOS}{r},${PUNTOS}${primero}:${PUNTOS}${ultimo}))'
        for k in range(1, len(COLS) + 1):
            letra = chr(64 + k)
            x = ws[f"{letra}{r}"]
            x.border = borde
            x.font = font(color="0000FF") if letra in (VALOR, MOVIL, HORARIO, WEB, NOTAS) else font()
            x.alignment = Alignment(wrap_text=True, vertical="top",
                                    horizontal="center" if letra in CENTRADAS else None)
            if letra in EDITABLES:
                x.fill = amarillo
        for letra in (VALOR, MOVIL, HORARIO, WEB):
            ws[f"{letra}{r}"].fill = azul
        ws[f"{MAPS}{r}"].font = font(color="0563C1", underline="single")
        if ws[f"{PRIM}{r}"].value:
            ws[f"{PRIM}{r}"].fill = verde
            ws[f"{PRIM}{r}"].font = font(bold=True, color="1F4E3D")
    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{RANK}{ultimo}"

    si_no = DataValidation(type="list", formula1='"Sí,No"', allow_blank=True)
    si_no_q = DataValidation(type="list", formula1='"Sí,No,?"', allow_blank=True)
    ws.add_data_validation(si_no)
    ws.add_data_validation(si_no_q)
    for letra in (QUEJA, LLAMADA, CADENA):
        si_no.add(f"{letra}{primero}:{letra}{ultimo}")
    si_no_q.add(f"{HORARIO}{primero}:{HORARIO}{ultimo}")
    rango = f"{PRIOR}{primero}:{PRIOR}{ultimo}"
    for texto, color in (("Alta", "F8C9C4"), ("Media", "FFE3A3"), ("Descartar", "DDDDDD")):
        ws.conditional_formatting.add(rango, CellIsRule(
            operator="equal", formula=[f'"{texto}"'], fill=PatternFill("solid", bgColor=color, fgColor=color)))

    g = wb.create_sheet("Cómo usarla")
    g.column_dimensions["A"].width = 30
    g.column_dimensions["B"].width = 95
    filas = [
        ("Qué es", "Hoja para comprobar a mano, en Google Maps, si estas clínicas pierden llamadas. Tú anotas lo que ves; el ranking se recalcula solo."),
        ("", ""),
        ("Colores", "Azul claro con texto azul = dato que preparé yo (puedes corregirlo). Amarillo = lo rellenas tú. Columnas R, S y T = fórmulas, no las toques."),
        ("", ""),
        ("Paso 1", "Empieza por las filas con ★ Primero. Haz clic en «Abrir en Maps»."),
        ("Paso 2", "Columna L: apunta el nº total de reseñas que ves en la ficha. Columna J: mira el horario; si cierra a mediodía o antes de las 18:00 entre semana, pon Sí (si no, No). Columna K: si en la ficha de Maps no hay botón de «Sitio web», escribe Sin web (suma 1 punto); si lo hay, pon la dirección."),
        ("Paso 3", "Abre las reseñas, ordénalas por «Más recientes» y usa la lupa para buscar: teléfono, llamar, cogen, contestan, comunica, cita. Lee solo las de los últimos 3 años."),
        ("Paso 4", "Si alguna reseña reciente se queja de no poder contactar por teléfono: columna M = Sí y pega la frase en N. No copies el nombre del autor."),
        ("Paso 5 (opcional, la señal más fiable)", "Llama a la clínica 3 veces en horas distintas (mañana, mediodía, tarde). Si no cogen en ninguna, columna O = Sí. Si es una cadena o franquicia, columna P = Sí y se descarta."),
        ("Paso 6", "Mira R (prioridad), S (puntos) y T (ranking). Ordena por T para ver las más potentes."),
        ("", ""),
        ("Cómo se puntúa", "Prioridad: Alta (queja en reseñas o llamada de prueba sin respuesta) 5 · Media (horario con huecos) 3 · Por verificar 1. Más: valor de la especialidad (H), móvil (I), +1 si hay más de 100 reseñas en Google y +1 si confirmas «Sin web». Máximo 10. Una cadena puntúa 0."),
        ("Ojo", "Los puntos iniciales salen solo de datos de directorio, sin reseñas de Google. Es un ranking de hipótesis: cambiará cuando anotes lo que ves."),
        ("", ""),
        ("Ejemplo de fila rellenada", "Clínica Dental Ejemplo · J: Sí · K: Sin web · L: 187 · M: Sí · N: «Llamé tres días seguidos y no cogen el teléfono» · O: No · P: No → Prioridad Alta, 5 + 2 + 0 + 1 + 1 = 9 puntos."),
    ]
    for i, (a, b) in enumerate(filas, 1):
        g.cell(i, 1, a).font = font(bold=True)
        g.cell(i, 2, b).font = font()
        for col in (1, 2):
            g.cell(i, col).alignment = Alignment(vertical="top", wrap_text=True)
    g.cell(3, 2).fill = azul

    wb.calculation.fullCalcOnLoad = True
    wb.save(salida)
    return n


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    print(f"{construir(leer(sys.argv[1]), sys.argv[2])} clínicas en {sys.argv[2]}")
